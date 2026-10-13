from __future__ import annotations

import io
import json
import logging
import re
from collections import Counter

import httpx
from docx import Document
from docx.shared import Pt
from fastapi import FastAPI, File, Form, HTTPException, UploadFile, Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic_settings import BaseSettings, SettingsConfigDict
from pypdf import PdfReader


class Settings(BaseSettings):
    model_config = SettingsConfigDict(env_file=".env", extra="ignore")
    cloudflare_account_id: str = ""
    cloudflare_api_token: str = ""
    cloudflare_ai_model: str = "@cf/meta/llama-3.1-8b-instruct"
    allowed_origins: str = "http://localhost:3000"


settings = Settings()
logger = logging.getLogger("resai")
app = FastAPI(title="ResAI API", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[origin.strip() for origin in settings.allowed_origins.split(",")],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

MAX_FILE_BYTES = 5 * 1024 * 1024
MIN_JD_CHARS = 80
STOP_WORDS = {
    # articles / prepositions / generic verbs
    "and", "the", "for", "with", "you", "your", "yours", "are", "our", "ours",
    "this", "that", "these", "those", "from", "will", "have", "has", "had", "into",
    "their", "theirs", "they", "them", "who", "whom", "which", "what", "when",
    "where", "how", "why", "not", "but", "all", "any", "can", "could", "should",
    "would", "may", "might", "must", "shall", "also", "very", "just", "about",
    "above", "after", "again", "against", "between", "both", "during", "each",
    "more", "most", "such", "than", "too", "own", "same", "using", "used",
    "use", "including", "include", "includes", "other", "others", "per", "via",
    "we", "us", "to", "of", "in", "on", "a", "an", "is", "it", "its", "as",
    "at", "by", "be", "been", "being", "was", "were", "or", "if", "so", "do",
    "does", "did", "done", "are", "am", "an", "no", "yes", "up", "out", "off",
    "over", "under", "then", "once", "here", "there", "s", "t", "re", "ll",
    "don", "will", "just", "with", "within", "without", "across", "along",
    "among", "around", "towards", "through", "upon", "onto", "per",
    "writing", "written", "collaborating", "collaboration", "problem",
    "solving", "solve", "ensure", "ensuring", "maintain", "maintaining",
    "support", "supporting", "develop", "developed", "developer", "developing",
    "design", "designed", "implement", "implemented", "implementing",
    "manage", "managed", "managing", "lead", "leading", "drive", "driving",
    # HR / JD filler that dominated top-30 and flattened scores to ~33%
    "work", "working", "role", "roles", "team", "teams", "years", "year",
    "experience", "experiences", "experienced", "skills", "skill", "required",
    "requirements", "requirement", "responsibilities", "responsibility",
    "qualifications", "qualification", "looking", "seeking", "hiring", "join",
    "joins", "joined", "opportunity", "opportunities", "position", "positions",
    "company", "business", "candidate", "candidates", "ideal", "great",
    "excellent", "strong", "good", "best", "top", "fast", "paced", "growing",
    "dynamic", "passion", "passionate", "motivated", "ability", "able",
    "help", "helps", "day", "daily", "etc", "eg", "ie", "plus", "highly",
    "proven", "solid", "deep", "broad", "diverse", "various", "multiple",
    "knowledge", "understanding", "background", "detail", "details",
    "attention", "communication", "collaborat", "analytical", "analysis",
    "reports", "stakeholder", "product", "project", "designing", "building",
    "create", "creates", "analyze", "generate", "needed", "must", "know",
    "ideal", "seniority", "level",
}


def extract_resume_text(content: bytes, filename: str) -> str:
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    text = ""
    try:
        if extension == "pdf":
            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
            if len(text.strip()) < 40:
                try:
                    from pdfminer.high_level import extract_text as pdfminer_extract

                    fallback = pdfminer_extract(io.BytesIO(content)) or ""
                    if len(fallback.strip()) > len(text.strip()):
                        text = fallback
                except Exception:
                    pass
        elif extension == "docx":
            document = Document(io.BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        else:
            raise HTTPException(400, "Upload a PDF or DOCX file.")
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(400, "We could not read that resume. Please upload a valid PDF or DOCX.") from error
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 40:
        raise HTTPException(
            400,
            "We couldn't extract text from this file — it may be a scanned/image PDF. Please export it as a text-based PDF or DOCX (e.g. File → Save as PDF / DOCX, or print to PDF) and try again.",
        )
    return text


def _clean_token(raw: str) -> str:
    # "team." -> "team", "next.js." -> "next.js", keep "c++"/"c#"
    word = raw.lower().strip()
    word = re.sub(r"[.,;:!?]+$", "", word)
    word = word.strip("()/\"'`")
    return word


def keywords(text: str, limit: int = 30) -> list[str]:
    raw_words = re.findall(r"[A-Za-z][A-Za-z0-9+#./-]{1,}", text.lower())
    cleaned: list[str] = []
    for raw in raw_words:
        word = _clean_token(raw)
        if len(word) < 2:
            continue
        if word in STOP_WORDS:
            continue
        if re.fullmatch(r"\d+", word or ""):
            continue
        cleaned.append(word)
    return [word for word, _ in Counter(cleaned).most_common(limit)]


def _term_in_text(term: str, text_lower: str) -> bool:
    # word-boundary match so "rest" != "forest", "bi" != "combine".
    # (?<!\w)...(?!\w) works for tech tokens like c++, c#, node.js too.
    try:
        return re.search(r"(?<!\w)" + re.escape(term) + r"(?!\w)", text_lower) is not None
    except re.error:
        return term in text_lower


def match_score(job_description: str, resume: str) -> int:
    jd_terms = set(keywords(job_description))
    resume_lower = resume.lower()
    if not jd_terms:
        return 0
    hits = sum(1 for term in jd_terms if _term_in_text(term, resume_lower))
    return round(100 * hits / len(jd_terms))


def llm_match_score(matched: list | None, missing: list | None, fallback: int) -> int:
    matched = matched or []
    missing = missing or []
    total = len(matched) + len(missing)
    if total == 0:
        return fallback
    return round(100 * len(matched) / total)


SYSTEM_PROMPT = """You are an expert resume editor. Create ATS-friendly resume content tailored to the job description, but never invent employers, education, dates, skills, achievements, qualifications, or metrics. Keep all facts grounded in the source resume. Return valid JSON only with these fields: tailored_resume (a polished plain-text resume, preserving known facts), matched_skills (array of strings), missing_keywords (array of strings), changes (array of concise strings)."""


async def tailor_with_llm(job_description: str, resume: str) -> dict:
    if not settings.cloudflare_account_id or not settings.cloudflare_api_token:
        raise HTTPException(503, "Cloudflare Workers AI is not configured. Add CLOUDFLARE_ACCOUNT_ID and CLOUDFLARE_API_TOKEN to backend/.env and restart the API.")
    payload = {
        "temperature": 0.25,
        "max_tokens": 4096,
        "response_format": {"type": "json_object"},
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": f"JOB DESCRIPTION:\n{job_description}\n\nSOURCE RESUME:\n{resume}"},
        ],
    }
    try:
        async with httpx.AsyncClient(timeout=60) as client:
            response = await client.post(
                "https://api.cloudflare.com/client/v4/accounts/"
                f"{settings.cloudflare_account_id}/ai/run/{settings.cloudflare_ai_model}",
                headers={
                    "Authorization": f"Bearer {settings.cloudflare_api_token}",
                    "Content-Type": "application/json",
                },
                json=payload,
            )
        response_data = response.json()
        if response.is_error:
            errors = response_data.get("errors", []) if isinstance(response_data, dict) else []
            message = "; ".join(str(item.get("message", item)) for item in errors) or "No error details returned"
            logger.warning("Cloudflare Workers AI returned HTTP %s: %s", response.status_code, message)
            raise HTTPException(502, f"Cloudflare Workers AI rejected the request: {message}")
        if not response_data.get("success", True):
            errors = response_data.get("errors", [])
            message = "; ".join(str(item.get("message", item)) for item in errors) or "No error details returned"
            logger.warning("Cloudflare Workers AI returned an unsuccessful response: %s", message)
            raise HTTPException(502, f"Cloudflare Workers AI rejected the request: {message}")
        result_payload = response_data.get("result", {})
        content = result_payload.get("response")
        if not isinstance(content, str):
            content = result_payload["choices"][0]["message"]["content"]
        # Some models occasionally wrap an otherwise-valid object in Markdown fences.
        content = content.strip()
        if content.startswith("```"):
            content = re.sub(r"^```(?:json)?\s*|\s*```$", "", content, flags=re.IGNORECASE).strip()
        result = json.loads(content)
    except HTTPException:
        raise
    except httpx.HTTPError as error:
        logger.exception("Cloudflare Workers AI network failure")
        raise HTTPException(502, "Could not reach Cloudflare Workers AI. Check your network and try again.") from error
    except (KeyError, ValueError, json.JSONDecodeError) as error:
        logger.exception("Cloudflare Workers AI returned an unreadable response")
        raise HTTPException(502, "Cloudflare Workers AI returned an unreadable response. Please try again.") from error
    if not isinstance(result.get("tailored_resume"), str):
        raise HTTPException(502, "The AI provider returned an invalid response. Please try again.")
    return result


@app.get("/health")
def health() -> dict:
    return {"status": "ok", "llm_configured": bool(settings.cloudflare_account_id and settings.cloudflare_api_token)}


@app.post("/api/customize")
async def customize_resume(
    job_description: str = Form(...), resume: UploadFile = File(...)
) -> dict:
    job_description = job_description.strip()
    if len(job_description) < MIN_JD_CHARS:
        raise HTTPException(400, f"Please paste at least {MIN_JD_CHARS} characters of the job description.")
    if not resume.filename:
        raise HTTPException(400, "Please select a resume file.")
    content = await resume.read()
    if not content:
        raise HTTPException(400, "The uploaded file is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(400, "Resume files must be 5 MB or smaller.")
    source_text = extract_resume_text(content, resume.filename)
    result = await tailor_with_llm(job_description, source_text)
    tailored_text = result["tailored_resume"].strip()
    matched = result.get("matched_skills", []) or []
    missing = result.get("missing_keywords", []) or []
    keyword_fallback = match_score(job_description, tailored_text)
    score = llm_match_score(matched, missing, keyword_fallback)
    return {
        "tailored_resume": tailored_text,
        "matched_skills": matched,
        "missing_keywords": missing,
        "changes": result.get("changes", []),
        "match_score": score,
        "source_filename": resume.filename,
    }


@app.post("/api/export/docx")
async def export_docx_endpoint(
    resume_text: str = Form(...),
    filename: str = Form("tailored-resume.docx"),
) -> Response:
    if not resume_text.strip():
        raise HTTPException(400, "Resume text is empty.")

    doc = Document()
    lines = resume_text.strip().split("\n")
    if lines:
        title_p = doc.add_paragraph()
        run = title_p.add_run(lines[0].strip())
        run.bold = True
        run.font.size = Pt(18)

        for line in lines[1:]:
            trimmed = line.strip()
            if not trimmed:
                doc.add_paragraph()
                continue
            if trimmed.startswith(("-", "*", "•")):
                doc.add_paragraph(trimmed.lstrip("-*• ").strip(), style="List Bullet")
            elif trimmed.isupper() and len(trimmed) < 40:
                doc.add_heading(trimmed, level=2)
            else:
                doc.add_paragraph(trimmed)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    clean_name = filename if filename.endswith(".docx") else f"{filename}.docx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{clean_name}"'},
    )

